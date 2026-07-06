import sys
import os
import subprocess

# Auto-install python-pptx if not present
try:
    import pptx
except ImportError:
    print("python-pptx library not found. Installing...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    import pptx

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    
    # Set slide dimensions to widescreen (16:9)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Theme Colors
    DARK_NAVY = RGBColor(15, 23, 42)      # #0f172a
    LIGHT_BG = RGBColor(248, 250, 252)    # #f8fafc
    WHITE = RGBColor(255, 255, 255)
    PRIMARY_BLUE = RGBColor(37, 99, 235)  # #2563eb
    TEXT_DARK = RGBColor(30, 41, 59)      # #1e293b
    TEXT_LIGHT = RGBColor(148, 163, 184)  # #94a3b8
    ACCENT_GREEN = RGBColor(16, 185, 129) # #10b981
    ACCENT_RED = RGBColor(239, 68, 68)    # #ef4444
    
    # Helper to apply background color
    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # Helper to add standard header to content slides
    def add_slide_header(slide, title_text):
        # Draw a subtle top accent bar
        top_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), Inches(13.333), Inches(0.15)
        )
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = PRIMARY_BLUE
        top_bar.line.fill.background()

        # Add Title Box
        title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.4), Inches(11.833), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Trebuchet MS"
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_BLUE

    # ----------------------------------------------------
    # SLIDE 1: Title Slide (Dark Theme)
    # ----------------------------------------------------
    slide_layout = prs.slide_layouts[6] # Blank slide
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, DARK_NAVY)
    
    # Add decorative colored shape
    accent_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.75), Inches(2.2), Inches(0.12), Inches(3.2)
    )
    accent_box.fill.solid()
    accent_box.fill.fore_color.rgb = PRIMARY_BLUE
    accent_box.line.fill.background()
    
    # Text Frame
    text_box = slide.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.0), Inches(3.5))
    tf = text_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    # Main Title
    p_title = tf.paragraphs[0]
    p_title.text = "PHISHGUARD"
    p_title.font.name = "Trebuchet MS"
    p_title.font.size = Pt(64)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    p_title.space_after = Pt(8)
    
    # Subtitle
    p_sub = tf.add_paragraph()
    p_sub.text = "A Hybrid Client-Server System for Real-Time Phishing & Cyber Crime Mitigation"
    p_sub.font.name = "Calibri"
    p_sub.font.size = Pt(22)
    p_sub.font.color.rgb = ACCENT_GREEN
    p_sub.space_after = Pt(45)
    
    # Team Details
    p_team = tf.add_paragraph()
    p_team.text = "Team Name: PhishGuard Solutions"
    p_team.font.name = "Calibri"
    p_team.font.size = Pt(18)
    p_team.font.bold = True
    p_team.font.color.rgb = WHITE
    
    p_members = tf.add_paragraph()
    p_members.text = "Team Members: [Member 1 (Developer), Member 2 (Designer/Researcher), Member 3]"
    p_members.font.name = "Calibri"
    p_members.font.size = Pt(15)
    p_members.font.color.rgb = TEXT_LIGHT
    
    # ----------------------------------------------------
    # SLIDE 2: Problem Statement (Light Theme)
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header(slide, "The Rising Cyber Crime Threat")
    
    # Left Column: Problem description
    left_col = slide.shapes.add_textbox(Inches(0.75), Inches(1.5), Inches(5.5), Inches(5.0))
    tf_left = left_col.text_frame
    tf_left.word_wrap = True
    tf_left.margin_left = tf_left.margin_top = tf_left.margin_right = tf_left.margin_bottom = 0
    
    p_prob_hdr = tf_left.paragraphs[0]
    p_prob_hdr.text = "The Phishing & Social Engineering Crisis"
    p_prob_hdr.font.name = "Trebuchet MS"
    p_prob_hdr.font.size = Pt(22)
    p_prob_hdr.font.bold = True
    p_prob_hdr.font.color.rgb = TEXT_DARK
    p_prob_hdr.space_after = Pt(15)
    
    bullet_points = [
        "Phishing accounts for over 80% of all reported security incidents worldwide, leading to massive financial and identity thefts.",
        "Zero-Day Domains: Cyber criminals deploy highly target-specific web links that stay active for less than 24 hours to bypass traditional blacklists (e.g., Google Safe Browsing).",
        "Quishing (QR Code Phishing): Physical flyers and QR codes are increasingly weaponized to redirect users to malicious login clones invisibly.",
        "Brand Impersonation: Sophisticated tricks like IDN Homograph attacks (fake letters mimicking real ones) fool even tech-savvy citizens."
    ]
    for pt in bullet_points:
        p = tf_left.add_paragraph()
        p.text = "• " + pt
        p.font.name = "Calibri"
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(12)
        
    # Right Column: Visual Callout Box
    right_bg = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.3)
    )
    right_bg.fill.solid()
    right_bg.fill.fore_color.rgb = RGBColor(254, 242, 242) # light red
    right_bg.line.color.rgb = ACCENT_RED
    right_bg.line.width = Pt(1.5)
    
    right_col = slide.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.1), Inches(3.7))
    tf_right = right_col.text_frame
    tf_right.word_wrap = True
    tf_right.margin_left = tf_right.margin_top = tf_right.margin_right = tf_right.margin_bottom = 0
    
    p_stat_title = tf_right.paragraphs[0]
    p_stat_title.text = "CRITICAL GAP IN CITIZEN SECURITY"
    p_stat_title.font.name = "Trebuchet MS"
    p_stat_title.font.size = Pt(18)
    p_stat_title.font.bold = True
    p_stat_title.font.color.rgb = ACCENT_RED
    p_stat_title.space_after = Pt(14)
    
    gaps = [
        "Inability of citizens to check suspicious links dynamically without visiting them.",
        "Emails look perfectly spoofed; email headers contain the proof (SPF/DKIM/DMARC) but are unreadable to regular users.",
        "Mobile security applications are often heavy, resource-intensive, and depend entirely on continuous cloud connections."
    ]
    for gap in gaps:
        p = tf_right.add_paragraph()
        p.text = "⚠️ " + gap
        p.font.name = "Calibri"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(14)

    # ----------------------------------------------------
    # SLIDE 3: Proposed Solution (Light Theme)
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header(slide, "Proposed Solution: PhishGuard")
    
    # Intro Subtitle
    intro_box = slide.shapes.add_textbox(Inches(0.75), Inches(1.3), Inches(11.833), Inches(0.6))
    p_intro = intro_box.text_frame.paragraphs[0]
    p_intro.text = "An intelligent, multi-layered anti-phishing ecosystem uniting local mobile AI with cloud-driven analysis."
    p_intro.font.name = "Calibri"
    p_intro.font.size = Pt(18)
    p_intro.font.italic = True
    p_intro.font.color.rgb = TEXT_DARK
    
    # 3 Pillars (Shapes/Cards)
    pillar_data = [
        {
            "title": "1. Local Mobile Engine",
            "desc": "A native Android app featuring an offline Machine Learning classifier (Logistic Regression) + heuristic scans to check links with zero latency and high battery efficiency.",
            "color": RGBColor(239, 246, 255), # light blue
            "border": PRIMARY_BLUE,
            "x": 0.75
        },
        {
            "title": "2. Server-Side Deep Scan",
            "desc": "A Node.js backend checking active DNS registration, validating SSL chains (detecting self-signed/expired certs), and auditing the page HTML layout for password credential theft scripts.",
            "color": RGBColor(240, 253, 250), # light teal
            "border": ACCENT_GREEN,
            "x": 4.85
        },
        {
            "title": "3. Security Sandbox",
            "desc": "An isolated, script-restricted preview sandbox that lets users view the visual layout of a suspicious website safely on their screens without executing any malicious payloads.",
            "color": RGBColor(255, 251, 235), # light yellow
            "border": RGBColor(217, 119, 6), # amber
            "x": 8.95
        }
    ]
    
    for pillar in pillar_data:
        # Background card shape
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(pillar["x"]), Inches(2.2), Inches(3.63), Inches(4.3)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = pillar["color"]
        card.line.color.rgb = pillar["border"]
        card.line.width = Pt(1.5)
        
        # Content box
        c_box = slide.shapes.add_textbox(
            Inches(pillar["x"] + 0.2), Inches(2.4), Inches(3.23), Inches(3.9)
        )
        tf_card = c_box.text_frame
        tf_card.word_wrap = True
        tf_card.margin_left = tf_card.margin_top = tf_card.margin_right = tf_card.margin_bottom = 0
        
        p_c_title = tf_card.paragraphs[0]
        p_c_title.text = pillar["title"]
        p_c_title.font.name = "Trebuchet MS"
        p_c_title.font.size = Pt(20)
        p_c_title.font.bold = True
        p_c_title.font.color.rgb = TEXT_DARK
        p_c_title.space_after = Pt(14)
        
        p_c_desc = tf_card.add_paragraph()
        p_c_desc.text = pillar["desc"]
        p_c_desc.font.name = "Calibri"
        p_c_desc.font.size = Pt(15)
        p_c_desc.font.color.rgb = TEXT_DARK
        p_c_desc.space_after = Pt(10)

    # ----------------------------------------------------
    # SLIDE 4: How the Solution Works (Light Theme)
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header(slide, "How the Solution Works: Detection Pipeline")
    
    # Draw horizontal pipeline steps
    steps = [
        {"num": "1", "name": "Input URL / QR / Email", "desc": "User enters URL, scans QR, or pastes raw email headers into PhishGuard."},
        {"num": "2", "name": "Local Heuristics Run", "desc": "Android client inspects Punycode (homographs), Levenshtein distance (typosquatting), and entropy."},
        {"num": "3", "name": "On-Device ML Predict", "desc": "Lightweight ML model runs locally to evaluate feature weights & output instant danger probability."},
        {"num": "4", "name": "API Deep Analysis", "desc": "Backend makes asynchronous DNS queries, SSL certificate analysis, and scans page HTML for login portals."}
    ]
    
    for i, step in enumerate(steps):
        left_pos = Inches(0.75 + (i * 2.95))
        
        # Step Number Bubble
        bubble = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            left_pos + Inches(1.1), Inches(1.7), Inches(0.75), Inches(0.75)
        )
        bubble.fill.solid()
        bubble.fill.fore_color.rgb = PRIMARY_BLUE
        bubble.line.fill.background()
        p_num = bubble.text_frame.paragraphs[0]
        p_num.text = step["num"]
        p_num.alignment = PP_ALIGN.CENTER
        p_num.font.name = "Trebuchet MS"
        p_num.font.size = Pt(20)
        p_num.font.bold = True
        p_num.font.color.rgb = WHITE
        
        # Card Background
        step_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left_pos, Inches(2.6), Inches(2.8), Inches(3.9)
        )
        step_card.fill.solid()
        step_card.fill.fore_color.rgb = WHITE
        step_card.line.color.rgb = RGBColor(226, 232, 240)
        
        # Card Text
        card_box = slide.shapes.add_textbox(
            left_pos + Inches(0.15), Inches(2.75), Inches(2.5), Inches(3.6)
        )
        tf_step = card_box.text_frame
        tf_step.word_wrap = True
        tf_step.margin_left = tf_step.margin_top = tf_step.margin_right = tf_step.margin_bottom = 0
        
        p_s_name = tf_step.paragraphs[0]
        p_s_name.text = step["name"]
        p_s_name.alignment = PP_ALIGN.CENTER
        p_s_name.font.name = "Trebuchet MS"
        p_s_name.font.size = Pt(16)
        p_s_name.font.bold = True
        p_s_name.font.color.rgb = TEXT_DARK
        p_s_name.space_after = Pt(12)
        
        p_s_desc = tf_step.add_paragraph()
        p_s_desc.text = step["desc"]
        p_s_desc.alignment = PP_ALIGN.CENTER
        p_s_desc.font.name = "Calibri"
        p_s_desc.font.size = Pt(13)
        p_s_desc.font.color.rgb = TEXT_DARK

    # ----------------------------------------------------
    # SLIDE 5: Major Components in the System (Light Theme)
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header(slide, "Major Components of PhishGuard")
    
    # 2x2 grid style representation
    comps = [
        ("Mobile App (Android / Kotlin Client)", 
         "Fully native Android app offering real-time URL interception, local heuristic scores, and a built-in camera reader to securely parse QR codes without executing scripts."),
        ("Dynamic Web Dashboard (Next.js PWA)", 
         "A highly responsive web application built with React, styled in Vanilla CSS, that offers manual link checking, security guide walkthroughs, and a responsive portal."),
        ("Deep Scan Node.js Backend API", 
         "A security API performing low-level certificate checks (Issuer, valid timeframe, fingerprint signatures), live DNS lookups, and checking HTML tags for spoof scripts."),
        ("Interactive Sandbox & Email Analyzer", 
         "Includes (1) A clean Safe Sandbox rendering iframe views securely with iframe scripting completely disabled, and (2) An Email Header parser diagnosing spoofed senders.")
    ]
    
    for idx, (title, desc) in enumerate(comps):
        row = idx // 2
        col = idx % 2
        left_pos = Inches(0.75 + (col * 6.0))
        top_pos = Inches(1.7 + (row * 2.7))
        
        # Draw component container
        comp_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            left_pos, top_pos, Inches(5.7), Inches(2.3)
        )
        comp_shape.fill.solid()
        comp_shape.fill.fore_color.rgb = WHITE
        comp_shape.line.color.rgb = RGBColor(226, 232, 240)
        comp_shape.line.width = Pt(1.5)
        
        # Indicator bar on left edge of the rectangle
        ind_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            left_pos, top_pos, Inches(0.12), Inches(2.3)
        )
        ind_bar.fill.solid()
        ind_bar.fill.fore_color.rgb = PRIMARY_BLUE if idx % 2 == 0 else ACCENT_GREEN
        ind_bar.line.fill.background()
        
        # Text block
        tb = slide.shapes.add_textbox(left_pos + Inches(0.3), top_pos + Inches(0.2), Inches(5.1), Inches(1.9))
        tf_comp = tb.text_frame
        tf_comp.word_wrap = True
        tf_comp.margin_left = tf_comp.margin_top = tf_comp.margin_right = tf_comp.margin_bottom = 0
        
        p_c_title = tf_comp.paragraphs[0]
        p_c_title.text = title
        p_c_title.font.name = "Trebuchet MS"
        p_c_title.font.size = Pt(18)
        p_c_title.font.bold = True
        p_c_title.font.color.rgb = TEXT_DARK
        p_c_title.space_after = Pt(6)
        
        p_c_desc = tf_comp.add_paragraph()
        p_c_desc.text = desc
        p_c_desc.font.name = "Calibri"
        p_c_desc.font.size = Pt(14)
        p_c_desc.font.color.rgb = TEXT_DARK

    # ----------------------------------------------------
    # SLIDE 6: Technical Tech Stack (Light Theme)
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header(slide, "System Architecture & Tech Stack")
    
    # Left Column: Technologies
    left_tb = slide.shapes.add_textbox(Inches(0.75), Inches(1.6), Inches(5.5), Inches(5.0))
    tf_left = left_tb.text_frame
    tf_left.word_wrap = True
    tf_left.margin_left = tf_left.margin_top = tf_left.margin_right = tf_left.margin_bottom = 0
    
    p_tech_hdr = tf_left.paragraphs[0]
    p_tech_hdr.text = "Technology Stack"
    p_tech_hdr.font.name = "Trebuchet MS"
    p_tech_hdr.font.size = Pt(22)
    p_tech_hdr.font.bold = True
    p_tech_hdr.font.color.rgb = TEXT_DARK
    p_tech_hdr.space_after = Pt(15)
    
    techs = [
        ("Mobile Client", "Native Kotlin, Android SDK, Local Heuristics Logic"),
        ("Frontend Web App", "Next.js 16 (React), Vanilla CSS, Tailwind CSS for dashboard structure"),
        ("Backend Services", "Next.js API Routes (Node.js/V8 execution context)"),
        ("Utilities & Scanner", "HTML5 QR-code parser, Web Cryptography API, DNS modules"),
        ("Model Infrastructure", "Python script for dataset extraction (`download_dataset.py`) & weights translation (`phishing_model_weights.json`)")
    ]
    for label, desc in techs:
        p = tf_left.add_paragraph()
        p.text = f"💻 {label}: "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = PRIMARY_BLUE
        
        p_desc = tf_left.add_paragraph()
        p_desc.text = "   " + desc
        p_desc.font.size = Pt(14)
        p_desc.font.color.rgb = TEXT_DARK
        p_desc.space_after = Pt(8)

    # Right Column: Key Algorithms
    right_tb = slide.shapes.add_textbox(Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tf_right = right_tb.text_frame
    tf_right.word_wrap = True
    tf_right.margin_left = tf_right.margin_top = tf_right.margin_right = tf_right.margin_bottom = 0
    
    p_algo_hdr = tf_right.paragraphs[0]
    p_algo_hdr.text = "Algorithms & Core Security Models"
    p_algo_hdr.font.name = "Trebuchet MS"
    p_algo_hdr.font.size = Pt(22)
    p_algo_hdr.font.bold = True
    p_algo_hdr.font.color.rgb = TEXT_DARK
    p_algo_hdr.space_after = Pt(15)
    
    algos = [
        ("Levenshtein Distance Metric", "Calculates string similarity to stop spoofed domains mimicking top-brand sites (PayPal, Google)."),
        ("Logistic Regression Probabilistic Model", "Processes URL length, special chars (@, _, -, .), subdomains, and protocols to evaluate a cumulative threat index."),
        ("Shannon Entropy Calculation", "Detects algorithmically generated, randomized random domains often deployed in spam/bot nets."),
        ("Punycode Decoding (`xn--`)", "Identifies IDN Homograph attacks where Cyrillic symbols hide within Latin domains to deceive end users.")
    ]
    for name, desc in algos:
        p = tf_right.add_paragraph()
        p.text = f"⚙️ {name}: "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = ACCENT_GREEN
        
        p_desc = tf_right.add_paragraph()
        p_desc.text = "   " + desc
        p_desc.font.size = Pt(14)
        p_desc.font.color.rgb = TEXT_DARK
        p_desc.space_after = Pt(8)

    # ----------------------------------------------------
    # SLIDE 7: Impact on Cybersecurity (Light Theme)
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, LIGHT_BG)
    add_slide_header(slide, "Impact on Cybersecurity & Society")
    
    # 4 grid blocks for impacts
    impacts = [
        ("Prevention Over Reaction", "Blocks malicious URLs before the user ever accesses them. Traditional browsers warn users after a site is flagged in a database; PhishGuard analyzes site structure instantly to block on Day-Zero."),
        ("Quishing (QR Phishing) Shield", "Protects citizens from physical scams. QR codes printed on fake parking slips or public posters are resolved locally, informing users of the underlying redirect risks before loading."),
        ("Accessible Email Security", "Democratizes complex header verification. Instead of looking at unreadable headers, users get simple ratings on whether the sender email addresses match verified SPF/DKIM profiles."),
        ("Education & Behavioral Shift", "Features gamified security guides and self-assessment quizzes that teach ordinary citizens how to identify phishing flags on their own, strengthening the human defense layer.")
    ]
    
    for idx, (title, desc) in enumerate(impacts):
        row = idx // 2
        col = idx % 2
        left_pos = Inches(0.75 + (col * 6.0))
        top_pos = Inches(1.8 + (row * 2.6))
        
        # Container shape
        container = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left_pos, top_pos, Inches(5.7), Inches(2.2)
        )
        container.fill.solid()
        container.fill.fore_color.rgb = WHITE
        container.line.color.rgb = RGBColor(226, 232, 240)
        
        tb = slide.shapes.add_textbox(left_pos + Inches(0.2), top_pos + Inches(0.2), Inches(5.3), Inches(1.8))
        tf_imp = tb.text_frame
        tf_imp.word_wrap = True
        tf_imp.margin_left = tf_imp.margin_top = tf_imp.margin_right = tf_imp.margin_bottom = 0
        
        p_t = tf_imp.paragraphs[0]
        p_t.text = "🎯 " + title
        p_t.font.name = "Trebuchet MS"
        p_t.font.size = Pt(17)
        p_t.font.bold = True
        p_t.font.color.rgb = PRIMARY_BLUE
        p_t.space_after = Pt(6)
        
        p_d = tf_imp.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Calibri"
        p_d.font.size = Pt(14)
        p_d.font.color.rgb = TEXT_DARK

    # ----------------------------------------------------
    # SLIDE 8: Demo & Conclusion (Dark Theme)
    # ----------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, DARK_NAVY)
    
    # Left Column: Demo highlights
    left_tb = slide.shapes.add_textbox(Inches(0.75), Inches(1.8), Inches(6.0), Inches(4.5))
    tf_left = left_tb.text_frame
    tf_left.word_wrap = True
    tf_left.margin_left = tf_left.margin_top = tf_left.margin_right = tf_left.margin_bottom = 0
    
    p_demo_hdr = tf_left.paragraphs[0]
    p_demo_hdr.text = "Working Prototype & Live Demo"
    p_demo_hdr.font.name = "Trebuchet MS"
    p_demo_hdr.font.size = Pt(26)
    p_demo_hdr.font.bold = True
    p_demo_hdr.font.color.rgb = WHITE
    p_demo_hdr.space_after = Pt(20)
    
    demo_bullets = [
        "Interactive URL Analyzer dashboard running at localhost:3000.",
        "Live QR Scanning interface using webcam inputs (HTML5 integration).",
        "Active Deep Scan API executing DNS resolutions and parsing HTML forms synchronously.",
        "Fully modular Kotlin class (PhishingEngine.kt) verifying URL components in the Android app."
    ]
    for db in demo_bullets:
        p = tf_left.add_paragraph()
        p.text = "⚡ " + db
        p.font.name = "Calibri"
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(14)

    # Right Column: Big Final Banner (Card)
    right_card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.2), Inches(1.8), Inches(5.3), Inches(4.2)
    )
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = PRIMARY_BLUE
    right_card.line.fill.background()
    
    right_tb = slide.shapes.add_textbox(Inches(7.5), Inches(2.2), Inches(4.7), Inches(3.4))
    tf_right = right_tb.text_frame
    tf_right.word_wrap = True
    tf_right.margin_left = tf_right.margin_top = tf_right.margin_right = tf_right.margin_bottom = 0
    
    p_end_title = tf_right.paragraphs[0]
    p_end_title.text = "PHISHGUARD"
    p_end_title.font.name = "Trebuchet MS"
    p_end_title.font.size = Pt(36)
    p_end_title.font.bold = True
    p_end_title.font.color.rgb = WHITE
    p_end_title.space_after = Pt(10)
    
    p_end_sub = tf_right.add_paragraph()
    p_end_sub.text = "Shielding Citizens from Cyber Scams"
    p_end_sub.font.name = "Calibri"
    p_end_sub.font.size = Pt(20)
    p_end_sub.font.color.rgb = ACCENT_GREEN
    p_end_sub.space_after = Pt(24)
    
    p_ty = tf_right.add_paragraph()
    p_ty.text = "Thank You! Questions?"
    p_ty.font.name = "Trebuchet MS"
    p_ty.font.size = Pt(22)
    p_ty.font.bold = True
    p_ty.font.color.rgb = WHITE

    # Save presentation
    output_filename = "ppt/PhishGuard_Presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation created successfully at {output_filename}!")

if __name__ == "__main__":
    create_presentation()
